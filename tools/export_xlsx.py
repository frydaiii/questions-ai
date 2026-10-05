"""Export all practice-bank questions to an Excel workbook.

Usage:
    python tools/export_xlsx.py
    python tools/export_xlsx.py --output exports/questions.xlsx
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import loader  # noqa: E402

DEFAULT_OUTPUT = loader.REPO_ROOT / "exports" / "aif-c01-questions.xlsx"
SHEET_NAME = "Questions"
HEADERS = [
    "Question ID",
    "Domain",
    "Task Statement",
    "Type",
    "Difficulty",
    "Case Study ID",
    "Select Count",
    "Stem",
    "Answer Choices / Matching Items",
    "Correct Answer",
    "Why Correct",
    "Why Incorrect",
    "Rule of Thumb",
    "Theory References",
    "Tags",
    "Source File",
    "Source Index",
]
COLUMN_WIDTHS = [
    16, 9, 16, 18, 12, 15, 14, 70, 70, 35, 70, 70, 55, 28, 30, 38, 14,
]


def clean_cell_value(value: Any) -> str:
    """Return an Excel-safe string within Excel's cell character limit."""
    text = "" if value is None else str(value)
    text = "".join(char for char in text if ord(char) >= 32 or char in "\n\t\r")
    return text[:32767]


def format_options(question: dict[str, Any]) -> str:
    """Render normal options or matching lists as readable multiline cell text."""
    if question.get("type") != "matching":
        return "\n".join(
            f"{option.get('key', '')}. {option.get('text', '')}"
            for option in question.get("options", [])
        )

    matching = question.get("matching", {})
    left = "\n".join(
        f"{item.get('key', '')}. {item.get('text', '')}"
        for item in matching.get("left", [])
    )
    right = "\n".join(
        f"{item.get('key', '')}. {item.get('text', '')}"
        for item in matching.get("right", [])
    )
    return f"Left items:\n{left}\n\nRight items:\n{right}"


def format_correct_answer(question: dict[str, Any]) -> str:
    """Render the answer key for all four question formats."""
    if question.get("type") == "matching":
        pairs = question.get("correct_pairs", {})
        return "\n".join(f"{left} -> {right}" for left, right in pairs.items())
    return ", ".join(question.get("correct", []))


def format_incorrect_explanations(question: dict[str, Any]) -> str:
    explanations = question.get("explanation", {}).get("why_incorrect", {})
    return "\n\n".join(
        f"{key}: {text}" for key, text in explanations.items()
    )


def question_row(record: loader.Record) -> list[str | int]:
    """Flatten a loaded question record into the worksheet column order."""
    question = record.data
    if question is None:
        raise ValueError(f"Cannot export unreadable question record: {record.location()}")

    explanation = question.get("explanation", {})
    try:
        source_file = record.file.relative_to(loader.REPO_ROOT)
    except ValueError:
        source_file = record.file

    return [
        clean_cell_value(question.get("id")),
        question.get("domain", ""),
        clean_cell_value(question.get("task_statement")),
        clean_cell_value(question.get("type")),
        clean_cell_value(question.get("difficulty")),
        clean_cell_value(question.get("case_study_id")),
        question.get("select_count", ""),
        clean_cell_value(question.get("stem")),
        clean_cell_value(format_options(question)),
        clean_cell_value(format_correct_answer(question)),
        clean_cell_value(explanation.get("why_correct")),
        clean_cell_value(format_incorrect_explanations(question)),
        clean_cell_value(explanation.get("rule_of_thumb")),
        clean_cell_value(", ".join(question.get("theory_refs", []))),
        clean_cell_value(", ".join(question.get("tags", []))),
        clean_cell_value(source_file),
        record.index,
    ]


def export_questions(output_path: Path) -> int:
    """Write all parsed question records to *output_path* and return their count."""
    records = loader.load_questions()
    invalid_records = [record for record in records if record.data is None]
    if invalid_records:
        locations = ", ".join(record.location() for record in invalid_records)
        raise ValueError(f"Cannot export because {len(invalid_records)} record(s) could not be read: {locations}")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = SHEET_NAME
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(records) + 1}"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for column, (header, width) in enumerate(zip(HEADERS, COLUMN_WIDTHS), start=1):
        cell = worksheet.cell(row=1, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.column_dimensions[get_column_letter(column)].width = width

    for row_index, record in enumerate(records, start=2):
        worksheet.append(question_row(record))
        for cell in worksheet[row_index]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.row_dimensions[1].height = 32
    worksheet.sheet_view.showGridLines = False
    workbook.properties.title = "AIF-C01 Practice Questions"
    workbook.properties.subject = "Exported practice question bank"
    workbook.properties.creator = "questions-ai export_xlsx"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return len(records)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export all practice questions to an XLSX workbook.")
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"Workbook destination (default: {DEFAULT_OUTPUT.relative_to(loader.REPO_ROOT)})",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output_path = args.output if args.output.is_absolute() else loader.REPO_ROOT / args.output
    try:
        count = export_questions(output_path)
    except ValueError as error:
        print(f"Export failed: {error}", file=sys.stderr)
        return 1

    try:
        display_path = output_path.relative_to(loader.REPO_ROOT)
    except ValueError:
        display_path = output_path
    print(f"Exported {count} question(s) to {display_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
